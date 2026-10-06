#!/usr/bin/env python3
"""Verify an exhibitor workbook without modifying it.
只读核验展商工作簿，不修改输入文件。
"""

from __future__ import annotations

import argparse
from copy import copy
import hashlib
import json
import re
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook


FORMULA_ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}


def sha256(path: Path) -> str:
    # Stream large files instead of loading the workbook binary into memory.
    # 分块读取文件，避免把整个工作簿二进制一次加载到内存。
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def style_signature(cell) -> tuple:
    # Style IDs are workbook-local; compare the actual formatting and protection.
    # 样式 ID 只在各自工作簿内有效；必须比较真实格式和单元格保护属性。
    return (
        copy(cell.font),
        copy(cell.fill),
        copy(cell.border),
        copy(cell.alignment),
        copy(cell.protection),
        cell.number_format,
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--sheet", required=True)
    parser.add_argument("--expected-rows", type=int, required=True, help="Data rows, excluding the header")
    parser.add_argument("--expected-columns", type=int, required=True)
    parser.add_argument("--baseline", type=Path)
    parser.add_argument("--immutable-columns", default="", help="Comma-separated Excel columns, for example A,B,C")
    parser.add_argument("--required-columns", default="", help="Comma-separated Excel columns that must be nonblank")
    parser.add_argument("--source-columns", default="", help="Comma-separated columns that must contain URLs or explicit status markers")
    parser.add_argument("--compare-all-styles", action="store_true", help="Compare every cell style against the baseline")
    parser.add_argument("--compare-layout", action="store_true", help="Compare merged ranges, freeze panes, row heights, and column widths")
    parser.add_argument("--forbidden-host", action="append", default=[])
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    # Comparisons need a baseline; silently ignoring these flags would mislead QA.
    # 对比校验必须提供基线；静默忽略这些参数会造成错误的验收结论。
    if args.expected_rows < 0 or args.expected_columns < 1:
        parser.error("expected rows must be nonnegative and expected columns positive")
    if (args.compare_all_styles or args.compare_layout or args.immutable_columns) and not args.baseline:
        parser.error("baseline is required for immutable/style/layout comparisons")

    workbook = load_workbook(args.workbook, data_only=False)
    if args.sheet not in workbook.sheetnames:
        raise SystemExit(f"sheet not found: {args.sheet}")
    sheet = workbook[args.sheet]
    errors: list[dict] = []

    actual_rows = sheet.max_row - 1
    actual_columns = sheet.max_column
    if actual_rows != args.expected_rows:
        errors.append({"type": "row_count", "expected": args.expected_rows, "actual": actual_rows})
    if actual_columns != args.expected_columns:
        errors.append({"type": "column_count", "expected": args.expected_columns, "actual": actual_columns})

    required = [value.strip().upper() for value in args.required_columns.split(",") if value.strip()]
    for column in required:
        for row in range(2, sheet.max_row + 1):
            if sheet[f"{column}{row}"].value in (None, ""):
                errors.append({"type": "blank_required", "cell": f"{column}{row}"})

    # Status text must not conceal a real hyperlink; validate that target as well.
    # 状态说明不能掩盖实际超链接；有超链接时仍需核验链接目标。
    forbidden = [re.compile(value, re.I) for value in args.forbidden_host]
    markers = {"Not publicly disclosed", "Pending verification", "Not applicable"}
    source_columns = [value.strip().upper() for value in args.source_columns.split(",") if value.strip()]
    for column in source_columns:
        for row in range(2, sheet.max_row + 1):
            cell = sheet[f"{column}{row}"]
            value = str(cell.value or "").strip()
            is_marker = any(value == marker or re.match(r"^" + re.escape(marker) + r"\s*[-–—:]", value) for marker in markers)
            if is_marker and not cell.hyperlink:
                continue
            text = cell.hyperlink.target if cell.hyperlink else value
            urls = re.findall(r"https?://[^\s;|]+", text)
            if not urls:
                errors.append({"type": "invalid_source", "cell": cell.coordinate, "value": value})
                continue
            for url in urls:
                try:
                    parsed = urlparse(url)
                    host = parsed.hostname
                    port = parsed.port
                except ValueError:
                    errors.append({"type": "invalid_source", "cell": cell.coordinate, "value": value})
                    continue
                if parsed.scheme not in {"http", "https"} or not host or parsed.username or parsed.password:
                    errors.append({"type": "invalid_source", "cell": cell.coordinate, "value": value})
                    continue
                if any(pattern.search(host) for pattern in forbidden):
                    errors.append({"type": "forbidden_source_host", "cell": cell.coordinate, "host": host})

    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value in FORMULA_ERRORS:
                errors.append({"type": "formula_error", "cell": cell.coordinate, "value": cell.value})

    # Check both formulas and saved cached values. This does not recalculate Excel.
    # 同时检查公式文本和已保存缓存；此步骤不负责重新计算 Excel 公式。
    cached_book = load_workbook(args.workbook, data_only=True)
    for row in cached_book[args.sheet].iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value in FORMULA_ERRORS:
                item = {"type": "formula_error", "cell": cell.coordinate, "value": cell.value}
                if item not in errors:
                    errors.append(item)

    # Detailed checks apply to the selected sheet; other sheets are name/order
    # guards, not a claim that their contents have been exhaustively compared.
    # 详细校验针对指定工作表；其他表只核名称和顺序，不声称已逐格完整对比。
    immutable = [value.strip().upper() for value in args.immutable_columns.split(",") if value.strip()]
    if args.baseline:
        baseline_book = load_workbook(args.baseline, data_only=False)
        if workbook.sheetnames != baseline_book.sheetnames:
            errors.append({"type": "worksheet_names_or_order_changed"})
        if args.sheet not in baseline_book.sheetnames:
            errors.append({"type": "baseline_sheet_missing", "sheet": args.sheet})
        else:
            baseline_sheet = baseline_book[args.sheet]
            if (sheet.max_row, sheet.max_column) != (baseline_sheet.max_row, baseline_sheet.max_column):
                errors.append({"type": "baseline_dimensions_changed"})
            for column in range(1, args.expected_columns + 1):
                if sheet.cell(1, column).value != baseline_sheet.cell(1, column).value:
                    errors.append({"type": "header_changed", "cell": sheet.cell(1, column).coordinate})
            for column in immutable:
                for row in range(1, min(sheet.max_row, baseline_sheet.max_row) + 1):
                    current = sheet[f"{column}{row}"]
                    original = baseline_sheet[f"{column}{row}"]
                    if current.value != original.value:
                        errors.append({"type": "immutable_value_changed", "cell": current.coordinate})
                    current_link = current.hyperlink.target if current.hyperlink else None
                    original_link = original.hyperlink.target if original.hyperlink else None
                    if current_link != original_link:
                        errors.append({"type": "immutable_hyperlink_changed", "cell": current.coordinate})
                    if style_signature(current) != style_signature(original):
                        errors.append({"type": "immutable_style_changed", "cell": current.coordinate})
            if args.compare_all_styles:
                for row in sheet.iter_rows():
                    for current in row:
                        if style_signature(current) != style_signature(baseline_sheet[current.coordinate]):
                            item = {"type": "cell_style_changed", "cell": current.coordinate}
                            errors.append(item)
            if args.compare_layout:
                # Hidden rows/columns and sheet protection are part of layout
                # preservation, even when no visible cell value has changed.
                # 隐藏行列和工作表保护也属于结构保真，即使可见数值未变化。
                if str(sheet.merged_cells) != str(baseline_sheet.merged_cells) or sheet.freeze_panes != baseline_sheet.freeze_panes:
                    errors.append({"type": "merges_or_freeze_panes_changed"})
                if sheet.sheet_state != baseline_sheet.sheet_state:
                    errors.append({"type": "sheet_visibility_changed"})
                if sheet.protection != baseline_sheet.protection:
                    errors.append({"type": "sheet_protection_changed"})
                for key in set(sheet.column_dimensions) | set(baseline_sheet.column_dimensions):
                    current = sheet.column_dimensions[key]
                    original = baseline_sheet.column_dimensions[key]
                    if (current.width, current.hidden, current.outlineLevel, current.collapsed) != (original.width, original.hidden, original.outlineLevel, original.collapsed):
                        errors.append({"type": "column_layout_changed", "column": key})
                for key in set(sheet.row_dimensions) | set(baseline_sheet.row_dimensions):
                    current = sheet.row_dimensions[key]
                    original = baseline_sheet.row_dimensions[key]
                    if (current.height, current.hidden, current.outlineLevel, current.collapsed) != (original.height, original.hidden, original.outlineLevel, original.collapsed):
                        errors.append({"type": "row_layout_changed", "row": key})

    # The report records a checksum of the unchanged input for traceable QA.
    # 报告记录未修改输入文件的校验值，供可追踪验收使用。
    report = {
        "schema": "exhibitor-workbook-validation-v1",
        "workbook": str(args.workbook.resolve()),
        "sha256": sha256(args.workbook),
        "sheet": args.sheet,
        "data_rows": actual_rows,
        "columns": actual_columns,
        "errors": errors,
        "status": "pass" if not errors else "fail",
    }
    text = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(text, encoding="utf-8")
    print(text, end="")
    return 0 if not errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
