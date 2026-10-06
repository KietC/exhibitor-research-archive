from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


# All inputs below are generated in temporary directories, never copied from
# production. Workbook verification is read-only; the tests also check that.
# 以下输入全部由临时目录合成，绝不复制生产数据；另核验工作簿校验确实只读。
class ManifestTests(unittest.TestCase):
    def test_generated_project_data_is_excluded(self):
        with tempfile.TemporaryDirectory(prefix="exhibitor-manifest-test-") as directory:
            root = Path(directory)
            (root / "source.py").write_text("print('example')\n", encoding="utf-8")
            for name in ("work", "data", "output", "evidence", "artifacts", "profile"):
                (root / name).mkdir()
                (root / name / "company.json").write_text('{"company":"Demo Alpha"}', encoding="utf-8")
            result = subprocess.run([sys.executable, str(ROOT / "scripts/build_manifest.py"), str(root)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            manifest = json.loads((root / "MANIFEST.sha256.json").read_text(encoding="utf-8"))
            self.assertEqual([row["path"] for row in manifest["files"]], ["source.py"])
            self.assertNotIn("root", manifest)

    def test_credentials_stop_packaging(self):
        with tempfile.TemporaryDirectory(prefix="exhibitor-credential-test-") as directory:
            root = Path(directory)
            (root / ".env").write_text("EXAMPLE=not-a-real-secret", encoding="utf-8")
            result = subprocess.run([sys.executable, str(ROOT / "scripts/build_manifest.py"), str(root)], capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            manifest = json.loads((root / "MANIFEST.sha256.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["rejected_credential_files"], [".env"])


class WorkbookTests(unittest.TestCase):
    def run_verifier(self, workbook, report, *extra):
        # Exercise the real CLI rather than a test-only reimplementation.
        # 直接测试真实命令行入口，而非另写一套测试专用逻辑。
        return subprocess.run(
            [sys.executable, str(ROOT / "scripts/verify_workbook.py"), str(workbook),
             "--sheet", "Exhibitors", "--expected-rows", "1", "--expected-columns", "1",
             "--report", str(report), *map(str, extra)],
            capture_output=True, text=True,
        )

    def test_malformed_sources_produce_failure_report(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory(prefix="exhibitor-source-url-test-") as directory:
            root = Path(directory)
            book = Workbook()
            book.active.title = "Exhibitors"
            book.active.append(["Source"])
            book.active.append(["https://[bad"])
            book.active.append(["https://:443"])
            book.active.append(["Pending verification"])
            book.active["A4"].hyperlink = "https://[bad"
            workbook, report = root / "test.xlsx", root / "report.json"
            book.save(workbook)
            result = subprocess.run([sys.executable, str(ROOT / "scripts/verify_workbook.py"), str(workbook), "--sheet", "Exhibitors", "--expected-rows", "3", "--expected-columns", "1", "--source-columns", "A", "--report", str(report)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 1, result.stderr)
            validation = json.loads(report.read_text(encoding="utf-8"))
            self.assertEqual(validation["status"], "fail")
            self.assertEqual({error["cell"] for error in validation["errors"]}, {"A2", "A3", "A4"})

    def test_style_comparison_is_semantic(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font

        spec = importlib.util.spec_from_file_location("verifier", ROOT / "scripts/verify_workbook.py")
        verifier = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(verifier)
        one, two = Workbook(), Workbook()
        one.active["A1"].font = Font(name="Arial", color="FF0000")
        two.active["A1"].font = Font(name="Arial", color="0000FF")
        self.assertNotEqual(verifier.style_signature(one.active["A1"]), verifier.style_signature(two.active["A1"]))

    def test_layout_changes_reported(self):
        from openpyxl import Workbook, load_workbook

        # Visibility, dimensions, outlines, and protection are separate hazards.
        # 可见性、尺寸、分组折叠与保护属性是彼此独立的结构风险。
        changes = {
            "hidden_row": (lambda sheet: setattr(sheet.row_dimensions[2], "hidden", True), "row_layout_changed"),
            "hidden_column": (lambda sheet: setattr(sheet.column_dimensions["A"], "hidden", True), "column_layout_changed"),
            "row_outline": (lambda sheet: setattr(sheet.row_dimensions[2], "outlineLevel", 1), "row_layout_changed"),
            "sheet_hidden": (lambda sheet: setattr(sheet, "sheet_state", "hidden"), "sheet_visibility_changed"),
            "sheet_protection": (lambda sheet: setattr(sheet.protection, "sheet", True), "sheet_protection_changed"),
        }
        for label, (mutate, expected_error) in changes.items():
            with self.subTest(label=label), tempfile.TemporaryDirectory(prefix="exhibitor-layout-test-") as directory:
                root = Path(directory)
                book = Workbook()
                sheet = book.active
                sheet.title = "Exhibitors"
                sheet.append(["Company"])
                sheet.append(["Demo Alpha"])
                book.create_sheet("Notes")
                baseline, current, report = root / "baseline.xlsx", root / "current.xlsx", root / "report.json"
                book.save(baseline)
                updated = load_workbook(baseline)
                mutate(updated["Exhibitors"])
                updated.save(current)
                result = self.run_verifier(current, report, "--baseline", baseline, "--compare-layout")
                self.assertEqual(result.returncode, 1, result.stderr)
                errors = json.loads(report.read_text(encoding="utf-8"))["errors"]
                self.assertIn(expected_error, {error["type"] for error in errors})

    def test_layout_unchanged_and_input_read_only(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory(prefix="exhibitor-read-only-test-") as directory:
            root = Path(directory)
            book = Workbook()
            sheet = book.active
            sheet.title = "Exhibitors"
            sheet.append(["Company"])
            sheet.append(["Demo Alpha"])
            sheet.row_dimensions[2].height = 18
            sheet.column_dimensions["A"].width = 24
            sheet.freeze_panes = "A2"
            sheet.protection.sheet = True
            workbook, report = root / "test.xlsx", root / "report.json"
            book.save(workbook)
            before = workbook.read_bytes()
            result = self.run_verifier(workbook, report, "--baseline", workbook, "--compare-layout", "--compare-all-styles", "--immutable-columns", "A")
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(workbook.read_bytes(), before)

    def test_cell_protection_style_reported(self):
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Protection

        with tempfile.TemporaryDirectory(prefix="exhibitor-cell-protection-test-") as directory:
            root = Path(directory)
            book = Workbook()
            book.active.title = "Exhibitors"
            book.active.append(["Company"])
            book.active.append(["Demo Alpha"])
            baseline, current, report = root / "baseline.xlsx", root / "current.xlsx", root / "report.json"
            book.save(baseline)
            updated = load_workbook(baseline)
            updated.active["A2"].protection = Protection(locked=False)
            updated.save(current)
            result = self.run_verifier(current, report, "--baseline", baseline, "--compare-all-styles")
            self.assertEqual(result.returncode, 1, result.stderr)
            errors = json.loads(report.read_text(encoding="utf-8"))["errors"]
            self.assertIn({"type": "cell_style_changed", "cell": "A2"}, errors)

    def test_immutable_hyperlink_changes_reported(self):
        from openpyxl import Workbook, load_workbook

        with tempfile.TemporaryDirectory(prefix="exhibitor-immutable-link-test-") as directory:
            root = Path(directory)
            book = Workbook()
            book.active.title = "Exhibitors"
            book.active.append(["Company"])
            book.active.append(["Demo Alpha"])
            book.active["A2"].hyperlink = "https://official.example/"
            baseline, current, report = root / "baseline.xlsx", root / "current.xlsx", root / "report.json"
            book.save(baseline)
            updated = load_workbook(baseline)
            updated.active["A2"].hyperlink = "https://unrelated.example/"
            updated.save(current)
            result = self.run_verifier(current, report, "--baseline", baseline, "--immutable-columns", "A")
            self.assertEqual(result.returncode, 1, result.stderr)
            errors = json.loads(report.read_text(encoding="utf-8"))["errors"]
            self.assertIn({"type": "immutable_hyperlink_changed", "cell": "A2"}, errors)

    def test_comparison_requires_baseline(self):
        # Invalid CLI configuration must fail before trying to read a workbook.
        # 命令行配置错误必须在尝试读取工作簿之前失败。
        result = self.run_verifier(Path("synthetic-missing.xlsx"), Path("synthetic-report.json"), "--compare-layout")
        self.assertEqual(result.returncode, 2)
        self.assertIn("baseline is required", result.stderr)


if __name__ == "__main__":
    unittest.main()
